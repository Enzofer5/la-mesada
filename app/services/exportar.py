"""Exportación de listados y reportes a CSV, Excel (xlsx) y PDF."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal

from flask import Response, current_app, has_request_context
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

VERDE = "3F6E3A"


def _nombre_archivo(titulo, extension):
    base = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode().lower()
    base = re.sub(r"[^a-z0-9]+", "_", base).strip("_") or "reporte"
    return f"{base}_{date.today():%Y%m%d}.{extension}"


def _texto_csv(valor):
    if valor is None:
        return ""
    if isinstance(valor, (float, Decimal)):
        return f"{valor:.2f}".replace(".", ",")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


def csv_respuesta(titulo, encabezados, filas):
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    escritor.writerow(encabezados)
    for fila in filas:
        escritor.writerow([_texto_csv(v) for v in fila])
    contenido = "﻿" + salida.getvalue()  # BOM para que Excel respete los acentos
    return Response(contenido, mimetype="text/csv; charset=utf-8", headers={
        "Content-Disposition": f"attachment; filename={_nombre_archivo(titulo, 'csv')}"})


def xlsx_bytes(titulo, encabezados, filas, subtitulo=None):
    libro = Workbook()
    hoja = libro.active
    hoja.title = re.sub(r"[\\/*?:\[\]]", "", titulo)[:31] or "Reporte"
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=14, color=VERDE)
    hoja.append([subtitulo or f"Generado el {datetime.now():%d/%m/%Y %H:%M}"])
    hoja.append([])
    hoja.append(encabezados)
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    for i, _ in enumerate(encabezados, start=1):
        letra = get_column_letter(i)
        ancho = max([len(str(encabezados[i - 1]))] + [len(_texto_csv(f[i - 1])) for f in filas[:500]]) + 2
        hoja.column_dimensions[letra].width = min(max(ancho, 10), 50)
        for celda in hoja[letra][4:]:
            if isinstance(celda.value, float):
                celda.number_format = "#,##0.00"
            elif isinstance(celda.value, (date, datetime)):
                celda.number_format = "DD/MM/YYYY"
    hoja.freeze_panes = "A5"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def xlsx_respuesta(titulo, encabezados, filas, subtitulo=None):
    return Response(xlsx_bytes(titulo, encabezados, filas, subtitulo),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={_nombre_archivo(titulo, 'xlsx')}"})


def pdf_bytes(titulo, encabezados, filas, subtitulo=None, empresa=None, totales=None):
    salida = io.BytesIO()
    horizontal = len(encabezados) > 5
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4) if horizontal else A4, leftMargin=12 * mm,
                            rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm, title=titulo)
    estilos = getSampleStyleSheet()
    estilos["Title"].textColor = colors.HexColor("#" + VERDE)
    celda_estilo = estilos["BodyText"].clone("celda", fontSize=8, leading=10)
    cabecera_estilo = celda_estilo.clone("cab", textColor=colors.white, fontName="Helvetica-Bold")
    elementos = []
    nombre_app = current_app.config.get("APP_NOMBRE", "La Mesada") if has_request_context() else "La Mesada"
    elementos.append(Paragraph(f"{empresa or nombre_app} · {titulo}", estilos["Title"]))
    elementos.append(Paragraph(subtitulo or "", estilos["Normal"]))
    elementos.append(Paragraph(f"Generado el {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]))
    elementos.append(Spacer(1, 6 * mm))

    def p(v, estilo=celda_estilo):
        texto = _texto_csv(v) if not isinstance(v, str) else v
        return Paragraph(texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), estilo)

    datos = [[p(e, cabecera_estilo) for e in encabezados]] + [[p(v) for v in fila] for fila in filas]
    if totales:
        datos.append([p(v, cabecera_estilo.clone("tot", textColor=colors.black)) for v in totales])
    if len(datos) == 1:
        datos.append([p("Sin datos para los filtros elegidos.")] + [p("")] * (len(encabezados) - 1))
    tabla = Table(datos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + VERDE)),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9cbb8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4ec")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if totales:
        estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e2e6d3")))
    tabla.setStyle(TableStyle(estilo))
    elementos.append(tabla)
    doc.build(elementos)
    return salida.getvalue()


def pdf_respuesta(titulo, encabezados, filas, subtitulo=None, empresa=None, totales=None):
    return Response(pdf_bytes(titulo, encabezados, filas, subtitulo, empresa, totales), mimetype="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename={_nombre_archivo(titulo, 'pdf')}"})


def respuesta(formato, titulo, encabezados, filas, subtitulo=None, totales=None):
    from .tenant import usuario_actual
    usuario = usuario_actual()
    empresa = usuario.empresa.nombre if usuario and usuario.empresa else None
    if formato == "csv":
        return csv_respuesta(titulo, encabezados, filas + ([totales] if totales else []))
    if formato == "xlsx":
        return xlsx_respuesta(titulo, encabezados, filas + ([totales] if totales else []), subtitulo)
    return pdf_respuesta(titulo, encabezados, filas, subtitulo, empresa, totales)
